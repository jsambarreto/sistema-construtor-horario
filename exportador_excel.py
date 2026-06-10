import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import hashlib

# Mapeamento humano dos 15 blocos lógicos que criámos na Fase 2
NOMES_BLOCOS = [
    "M1 (07:30-08:15)", "M2 (08:15-09:00)", "M3 (09:15-10:00)", 
    "M4 (10:00-10:45)", "M5 (10:45-11:30)", "M6 (11:30-12:15)",
    "T1 (13:15-14:00)", "T2 (14:00-14:45)", "T3 (15:00-15:45)", 
    "T4 (15:45-16:30)", "T5 (16:30-17:15)",
    "N1 (18:30-19:20)", "N2 (19:20-20:10)", "N3 (20:20-21:10)", "N4 (21:10-22:00)"
]

DIAS_SEMANA = ['SEG', 'TER', 'QUA', 'QUI', 'SEX']

def gerar_cor_hex(texto):
    """
    Gera uma cor pastel única e consistente baseada no nome da disciplina ou turma,
    para que, por exemplo, 'Matemática' tenha sempre a mesma cor na grelha.
    """
    hash_obj = hashlib.md5(texto.encode('utf-8'))
    hexa = hash_obj.hexdigest()[:6]
    # Clarear a cor para tom pastel misturando com branco (FF)
    r = (int(hexa[0:2], 16) + 255) // 2
    g = (int(hexa[2:4], 16) + 255) // 2
    b = (int(hexa[4:6], 16) + 255) // 2
    return f"{r:02X}{g:02X}{b:02X}"

def configurar_planilha(wb, nome_aba):
    """Cria uma aba formatada com cabeçalhos de dias e horários."""
    ws = wb.create_sheet(title=nome_aba[:31]) # Limite do Excel são 31 caracteres
    
    # Estilos Base
    fonte_cabecalho = Font(bold=True, color="FFFFFF")
    fundo_cabecalho = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    alinhamento_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), 
                        top=Side(style='thin'), bottom=Side(style='thin'))

    # Configurar Cabeçalhos (Dias da Semana)
    ws.cell(row=1, column=1, value="HORÁRIO").font = fonte_cabecalho
    ws.cell(row=1, column=1).fill = fundo_cabecalho
    ws.cell(row=1, column=1).alignment = alinhamento_centro
    ws.column_dimensions['A'].width = 18

    for col_idx, dia in enumerate(DIAS_SEMANA, start=2):
        celula = ws.cell(row=1, column=col_idx, value=dia)
        celula.font = fonte_cabecalho
        celula.fill = fundo_cabecalho
        celula.alignment = alinhamento_centro
        celula.border = borda_fina
        # Aumentar largura para caber "Disciplina \n Professor"
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 25 

    # Configurar Linhas (Blocos de Horário)
    for row_idx, bloco in enumerate(NOMES_BLOCOS, start=2):
        celula = ws.cell(row=row_idx, column=1, value=bloco)
        celula.alignment = alinhamento_centro
        celula.border = borda_fina
        # Diferenciar turnos visualmente na primeira coluna
        if "M" in bloco: cor_turno = "D9E1F2"
        elif "T" in bloco: cor_turno = "FFF2CC"
        else: cor_turno = "E2EFDA"
        celula.fill = PatternFill(start_color=cor_turno, end_color=cor_turno, fill_type="solid")

    return ws, alinhamento_centro, borda_fina

def extrair_dados_solver(solver, alocacoes):
    """Lê o solver resolvido e cria dicionários estruturados para Turmas e Docentes."""
    grade_turmas = {}   # grade_turmas[turma][dia][bloco] = "Disc \n Docente"
    grade_docentes = {} # grade_docentes[docente][dia][bloco] = "Turma \n Disc"
    
    for (turma, disc, doc, dia_idx, bloco), var in alocacoes.items():
        if solver.Value(var) == 1:
            # Inicializar estruturas se não existirem
            if turma not in grade_turmas: grade_turmas[turma] = {d: {} for d in range(5)}
            if doc not in grade_docentes: grade_docentes[doc] = {d: {} for d in range(5)}
            
            # Guardar alocação
            grade_turmas[turma][dia_idx][bloco] = f"{disc}\n({doc})"
            grade_docentes[doc][dia_idx][bloco] = f"{turma}\n{disc}"
            
    return grade_turmas, grade_docentes

# CORREÇÃO: Nome e assinatura ajustados para combinar com o main.py
def exportar_para_excel(solver, alocacoes, dias_semana, docentes, turmas, caminho_saida_turmas="Horarios_Turmas_2026.xlsx", caminho_saida_docentes="Horarios_Docentes_2026.xlsx"):
    print("\n[+] A iniciar geração dos ficheiros Excel...")
    grade_turmas, grade_docentes = extrair_dados_solver(solver, alocacoes)
    
    # ---------------------------------------------------------
    # 1. EXCEL DAS TURMAS
    # ---------------------------------------------------------
    wb_turmas = openpyxl.Workbook()
    wb_turmas.remove(wb_turmas.active) # Remove aba padrão vazia
    
    for turma, dados_dias in grade_turmas.items():
        ws, alinhamento, borda = configurar_planilha(wb_turmas, turma)
        
        for dia_idx, blocos in dados_dias.items():
            for bloco_idx, conteudo in blocos.items():
                # No excel, a Linha 1 é cabeçalho, logo bloco 0 -> linha 2
                # A Coluna A (1) é o horário, logo dia_idx 0 (SEG) -> coluna 2
                linha = bloco_idx + 2
                coluna = dia_idx + 2
                
                celula = ws.cell(row=linha, column=coluna, value=conteudo)
                celula.alignment = alinhamento
                celula.border = borda
                
                # Pintar a célula com base na disciplina (para consistência visual)
                nome_disc = conteudo.split('\n')[0]
                cor = gerar_cor_hex(nome_disc)
                celula.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
                
    wb_turmas.save(caminho_saida_turmas)
    print(f"[SUCESSO] Ficheiro gerado: {caminho_saida_turmas}")
    
    # ---------------------------------------------------------
    # 2. EXCEL DOS DOCENTES
    # ---------------------------------------------------------
    wb_docentes = openpyxl.Workbook()
    wb_docentes.remove(wb_docentes.active)
    
    # Ordenar docentes alfabeticamente para facilitar a procura
    for docente in sorted(grade_docentes.keys()):
        dados_dias = grade_docentes[docente]
        ws, alinhamento, borda = configurar_planilha(wb_docentes, docente)
        
        for dia_idx, blocos in dados_dias.items():
            for bloco_idx, conteudo in blocos.items():
                linha = bloco_idx + 2
                coluna = dia_idx + 2
                
                celula = ws.cell(row=linha, column=coluna, value=conteudo)
                celula.alignment = alinhamento
                celula.border = borda
                
                # Pintar a célula com base na Turma
                nome_turma = conteudo.split('\n')[0]
                cor = gerar_cor_hex(nome_turma)
                celula.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
                
    wb_docentes.save(caminho_saida_docentes)
    print(f"[SUCESSO] Ficheiro gerado: {caminho_saida_docentes}")